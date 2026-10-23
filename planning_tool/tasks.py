from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from planning_tool.models import WorkPackage

TIME_RE = re.compile(r"^\s*(\d+)\s*h?\s*$", re.IGNORECASE)


def parse_time_cell(value: object) -> int:
    if value is None:
        raise ValueError("missing time")
    if isinstance(value, (int, float)):
        h = int(value)
        if h <= 0:
            raise ValueError(f"invalid hours: {value}")
        return h
    text = str(value).strip()
    m = TIME_RE.match(text)
    if not m:
        raise ValueError(f"invalid time format: {text!r} (expected e.g. 2h or 16)")
    return int(m.group(1))


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _merged_cell_value(ws: Worksheet, row: int, col: int) -> object:
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= col <= merged.max_col
        ):
            return ws.cell(row=merged.min_row, column=merged.min_col).value
    return None


def _rows_from_xlsx(path: Path) -> list[tuple[str, str, str, str]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    rows: list[tuple[str, str, str, str]] = []
    for row_idx in range(1, ws.max_row + 1):
        a = _cell_str(_merged_cell_value(ws, row_idx, 1))
        b = _cell_str(_merged_cell_value(ws, row_idx, 2))
        c = _cell_str(_merged_cell_value(ws, row_idx, 3))
        d = _cell_str(_merged_cell_value(ws, row_idx, 4))
        if not any((a, b, c, d)):
            continue
        rows.append((a, b, c, d))
    wb.close()
    return rows


def _rows_from_sheet(path: Path) -> list[tuple[str, str, str, str]]:
    suffix = path.suffix.lower()
    rows: list[tuple[str, str, str, str]] = []

    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for raw in reader:
                if not raw or all(not _cell_str(c) for c in raw):
                    continue
                a = _cell_str(raw[0]) if len(raw) > 0 else ""
                b = _cell_str(raw[1]) if len(raw) > 1 else ""
                c = _cell_str(raw[2]) if len(raw) > 2 else ""
                d = _cell_str(raw[3]) if len(raw) > 3 else ""
                rows.append((a, b, c, d))
        return rows

    if suffix in (".xlsx", ".xlsm"):
        return _rows_from_xlsx(path)

    raise ValueError(f"unsupported task file type: {path.suffix} (use .csv or .xlsx)")


def _is_header_row(milestone: str, prefix: str, name: str, time_raw: str) -> bool:
    joined = " ".join(x.lower() for x in (milestone, prefix, name, time_raw))
    if "work package" in joined or time_raw.lower() == "time":
        return True
    if name.lower() in ("work packages", "work package"):
        return True
    return False


def load_work_packages(path: Path) -> list[WorkPackage]:
    rows = _rows_from_sheet(path)
    packages: list[WorkPackage] = []
    last_prefix = ""
    idx = 0

    for _milestone, prefix_raw, name, time_raw in rows:
        if not packages and _is_header_row(_milestone, prefix_raw, name, time_raw):
            continue
        # B 列：合并单元格或逐行不同前缀均可；有值则更新，空则沿用上一行
        if prefix_raw:
            last_prefix = prefix_raw
        prefix = last_prefix
        if not name:
            continue
        if not prefix:
            raise ValueError(f"row with work package {name!r} has no Prefix (column B)")
        hours = parse_time_cell(time_raw)
        idx += 1
        packages.append(
            WorkPackage(index=idx, prefix=prefix, name=name, hours=hours)
        )

    if not packages:
        raise ValueError("no work packages found in task file")
    return packages


def distinct_prefixes(packages: list[WorkPackage]) -> list[str]:
    seen: list[str] = []
    for p in packages:
        if p.prefix not in seen:
            seen.append(p.prefix)
    return seen


def primary_prefix(packages: list[WorkPackage]) -> str:
    prefs = distinct_prefixes(packages)
    if len(prefs) == 1:
        return prefs[0]
    if not prefs:
        return ""
    return ", ".join(prefs)
